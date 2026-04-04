import openpyxl
import json
from pathlib import Path

excel_path = Path("/Users/prabhathvinay/Documents/ai-badge-classification-platform/prabhath/data/master/NJIT_Badge_Classification_FACT-SHEET.xlsx")
json_path = Path("/Users/prabhathvinay/Documents/ai-badge-classification-platform/prabhath/data/master/master_badges.json")

print(f"Reading: {excel_path}")
print(f"Output: {json_path}")

wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb.active

# Get headers from row 3
headers = [str(c).strip() if c else None for c in list(ws.iter_rows(min_row=3, max_row=3, values_only=True))[0]]
print(f"Found {len(headers)} columns")

# Find key columns by name
name_col = None
cat_col = None
type_col = None
level_col = None

for i, h in enumerate(headers):
    if not h:
        continue
    hl = h.lower()
    if "badge" in hl and "title" in hl:
        name_col = i
    elif "stage 1" in hl or ("category" in hl and "stage" in hl):
        cat_col = i
    elif "stage 2" in hl or ("type" in hl and "stage" in hl):
        type_col = i
    elif "stage 3" in hl or ("level" in hl and "stage" in hl):
        level_col = i

print(f"Columns: name={name_col}, cat={cat_col}, type={type_col}, level={level_col}")

badges = []
for row_idx, row in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
    if not row or not row[0]:
        continue

    try:
        name = row[name_col] if name_col is not None else None
        cat = row[cat_col] if cat_col is not None else None
        typ = row[type_col] if type_col is not None else None
        lvl = row[level_col] if level_col is not None else None

        if name and cat:
            badge = {
                "badge_name": str(name),
                "expected_category": str(cat) if cat else None,
                "expected_type": str(typ) if typ else None,
                "expected_level": str(lvl) if lvl else None,
                "source_format": "master_dataset"
            }
            badges.append(badge)
    except Exception as e:
        print(f"Error at row {row_idx}: {e}")

print(f"\nExtracted {len(badges)} badges")

# Write JSON
json_path.parent.mkdir(parents=True, exist_ok=True)
with open(json_path, "w") as f:
    json.dump(badges, f, indent=2)

print(f"Saved to {json_path}")

# Summary
categories = {}
types = {}
levels = {}
for b in badges:
    cat = b.get("expected_category", "Unknown")
    typ = b.get("expected_type", "Unknown")
    lvl = b.get("expected_level", "Unknown")
    categories[cat] = categories.get(cat, 0) + 1
    types[typ] = types.get(typ, 0) + 1
    levels[lvl] = levels.get(lvl, 0) + 1

print(f"\nCategories: {dict(categories)}")
print(f"Types: {dict(types)}")
print(f"Levels: {dict(levels)}")

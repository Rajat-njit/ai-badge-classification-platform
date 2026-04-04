import openpyxl
import json
from pathlib import Path

excel_path = Path("/Users/prabhathvinay/Documents/ai-badge-classification-platform/prabhath/data/master/NJIT_Badge_Classification_FACT-SHEET.xlsx")
json_path = Path("/Users/prabhathvinay/Documents/ai-badge-classification-platform/prabhath/data/master/master_badges.json")

print(f"Reading: {excel_path}")

wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb.active

# Get headers from row 3
headers = [str(c).strip() if c else None for c in list(ws.iter_rows(min_row=3, max_row=3, values_only=True))[0]]
print(f"Found {len(headers)} columns")

# Map column indices
name_col = None
cat_col = None
type_col = None
level_col = None
issuer_col = None
ach_type_col = None
assessment_type_col = None
assessment_eval_col = None
evidence_col = None
audience_col = None
bloom_col = None
pathway_col = None

for i, h in enumerate(headers):
    if not h:
        continue
    hl = h.lower()
    if "badge" in hl and "title" in hl:
        name_col = i
    elif "stage 1" in hl or ("category" in hl and "stage" in hl):
        cat_col = i
    elif "stage 2" in hl or ("type" in hl and "stage" in hl):
        type_col = i
    elif "stage 3" in hl or ("level" in hl and "stage" in hl):
        level_col = i
    elif "issuer" in hl:
        issuer_col = i
    elif "achievement type" in hl:
        ach_type_col = i
    elif "assessment type" in hl and "eval" not in hl:
        assessment_type_col = i
    elif "assessment evaluator" in hl:
        assessment_eval_col = i
    elif "evidence type" in hl:
        evidence_col = i
    elif "audience" in hl:
        audience_col = i
    elif "bloom" in hl:
        bloom_col = i
    elif "pathway position" in hl:
        pathway_col = i

print(f"Mapped: name={name_col}, issuer={issuer_col}, cat={cat_col}, type={type_col}, level={level_col}")
print(f"        audience={audience_col}, evidence={evidence_col}, bloom={bloom_col}, assessment={assessment_type_col}")

badges = []
for row_idx, row in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
    if not row or not row[0]:
        continue

    try:
        name = row[name_col] if name_col is not None else None
        cat = row[cat_col] if cat_col is not None else None

        if not name or not cat:
            continue

        badge = {
            "badge_name": str(name),
            "expected_category": str(cat),
            "source_format": "master_dataset",
            "source_row": row_idx
        }

        # Expected type/level
        if type_col is not None and row[type_col]:
            badge["expected_type"] = str(row[type_col])
        if level_col is not None and row[level_col]:
            badge["expected_level"] = str(row[level_col])

        # Issuer
        if issuer_col is not None and row[issuer_col]:
            badge["issuing_department"] = str(row[issuer_col])
        else:
            badge["issuing_department"] = "NJIT"

        # Description
        badge["description"] = f"NJIT Digital Badge - {badge['badge_name']}"

        # Audience mapping
        if audience_col is not None and row[audience_col]:
            aud_raw = str(row[audience_col]).lower()
            if "faculty" in aud_raw and "staff" in aud_raw:
                badge["audience"] = ["faculty", "staff"]
            elif "faculty" in aud_raw:
                badge["audience"] = ["faculty"]
            elif "staff" in aud_raw:
                badge["audience"] = ["staff"]
            elif "external" in aud_raw or "professional" in aud_raw:
                badge["audience"] = ["external", "professional"]
            elif "student" in aud_raw:
                badge["audience"] = ["student"]
            else:
                badge["audience"] = ["student"]
        else:
            # Infer from category
            cat_lower = str(cat).lower()
            if "faculty" in cat_lower or "staff" in cat_lower:
                badge["audience"] = ["faculty", "staff"]
            elif "professional" in cat_lower or "continuing" in cat_lower:
                badge["audience"] = ["external", "professional"]
            else:
                badge["audience"] = ["student"]

        # Context from category
        cat_lower = str(cat).lower()
        if "faculty" in cat_lower or "staff" in cat_lower:
            badge["context"] = "employee_development"
        elif "professional" in cat_lower or "continuing" in cat_lower:
            badge["context"] = "professional_ed"
        elif "co-curricular" in cat_lower:
            badge["context"] = "co_curricular"
        else:
            badge["context"] = "academic"

        # Assessment info
        has_assessment = True
        assessment_desc = None

        if assessment_type_col is not None and row[assessment_type_col]:
            assessment_raw = str(row[assessment_type_col]).lower()
            assessment_desc = str(row[assessment_type_col])
            if assessment_raw in ["attendance", "attendance only", "none", "n/a", ""]:
                has_assessment = False
                badge["assessment_type"] = None
            elif "module" in assessment_raw or "quiz" in assessment_raw or "exam" in assessment_raw:
                badge["assessment_type"] = "auto_graded"
            elif "rubric" in assessment_raw or "panel" in assessment_raw or "judge" in assessment_raw:
                badge["assessment_type"] = "expert_multi"
            elif "expert" in assessment_raw or "evaluator" in assessment_raw or "demo" in assessment_raw:
                badge["assessment_type"] = "expert_skill"
            else:
                badge["assessment_type"] = "auto_graded"
        else:
            badge["assessment_type"] = "auto_graded"

        badge["has_assessment"] = has_assessment
        badge["assessment_description"] = assessment_desc

        # Completion criteria
        badge["completion_criteria"] = "Complete all requirements"

        # Evidence
        evidence = []
        if evidence_col is not None and row[evidence_col]:
            ev_raw = str(row[evidence_col]).lower()
            if "attendance" in ev_raw:
                evidence.append("attendance_record")
            if "module" in ev_raw or "quiz" in ev_raw:
                evidence.append("quiz_result")
            if "artifact" in ev_raw or "project" in ev_raw or "product" in ev_raw:
                evidence.append("project_artifact")
            if "rubric" in ev_raw:
                evidence.append("expert_rated_rubric")
            if "portfolio" in ev_raw:
                evidence.append("portfolio")
            if "presentation" in ev_raw or "pitch" in ev_raw:
                evidence.append("presentation")
            if "review" in ev_raw or "eval" in ev_raw:
                evidence.append("expert_review")
            if "application" in ev_raw:
                evidence.append("assignment_submission")
            if "plan" in ev_raw:
                evidence.append("project_artifact")
            if "360" in ev_raw or "peer" in ev_raw:
                evidence.append("360_feedback")
            if "supervisor" in ev_raw:
                evidence.append("supervisor_evaluation")

        if not evidence:
            evidence = ["attendance_record"]
        badge["evidence_required"] = evidence

        # Bloom levels
        bloom = []
        if bloom_col is not None and row[bloom_col]:
            bloom_raw = str(row[bloom_col]).lower()
            if "remember" in bloom_raw:
                bloom.append("remember")
            if "understand" in bloom_raw or "understanding" in bloom_raw:
                bloom.append("understand")
            if "apply" in bloom_raw or "application" in bloom_raw:
                bloom.append("apply")
            if "analyze" in bloom_raw or "analysis" in bloom_raw:
                bloom.append("analyze")
            if "evaluate" in bloom_raw or "evaluation" in bloom_raw:
                bloom.append("evaluate")
            if "create" in bloom_raw:
                bloom.append("create")

        if not bloom:
            # Default based on type
            exp_type = badge.get("expected_type", "").lower()
            if "souvenir" in exp_type:
                bloom = []
            elif "achievement" in exp_type:
                bloom = ["understand"]
            elif "skill" in exp_type:
                bloom = ["apply"]
            else:
                bloom = ["understand"]

        badge["bloom_levels"] = bloom

        # Is terminal based on pathway position or level
        is_terminal = False
        if pathway_col is not None and row[pathway_col]:
            pos = str(row[pathway_col]).lower()
            if "terminal" in pos or "final" in pos or "capstone" in pos:
                is_terminal = True

        exp_level = badge.get("expected_level", "").lower()
        if "terminal" in exp_level or "mastery" in exp_level or "exemplary" in exp_level or "capstone" in exp_level:
            is_terminal = True

        badge["is_terminal"] = is_terminal
        badge["prerequisites"] = []
        badge["hours_to_complete"] = 10

        badges.append(badge)

    except Exception as e:
        print(f"Error at row {row_idx}: {e}")

print(f"\nExtracted {len(badges)} badges")

# Write JSON
json_path.parent.mkdir(parents=True, exist_ok=True)
with open(json_path, "w") as f:
    json.dump(badges, f, indent=2)

print(f"Saved to {json_path}")

# Summary
categories = {}
types = {}
levels = {}
for b in badges:
    cat = b.get("expected_category", "Unknown")
    typ = b.get("expected_type", "Unknown")
    lvl = b.get("expected_level", "Unknown")
    categories[cat] = categories.get(cat, 0) + 1
    types[typ] = types.get(typ, 0) + 1
    levels[lvl] = levels.get(lvl, 0) + 1

print(f"\nCategories: {dict(categories)}")
print(f"Types: {dict(types)}")
print(f"Levels: {dict(levels)}")
